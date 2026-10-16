#!/usr/bin/env python3
"""Raport Yahoo Finance: Analyst Recommendation + Top Analysts.

Exemplu:
    python yahoo_analyst_report.py AAPL
"""

from __future__ import annotations

import argparse
import html
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
import yfinance as yf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rich import box
from rich.console import Console
from rich.panel import Panel
from rich.table import Table
from rich.text import Text


# Cardurile „Analyst Insights”, inclusiv „Latest Rating”, sunt în pagina
# principală a simbolului. URL-ul vechi /analysts/ redirecționează frecvent aici.
ANALYSTS_URL = "https://finance.yahoo.com/quote/{ticker}/"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36",
    "Accept": "application/json,text/html,application/xhtml+xml",
}
SESSION = requests.Session()
console = Console()
ANALYSTS_PAGE_CACHE: dict[str, str] = {}


class YahooDataError(RuntimeError):
    """Datele cerute nu au putut fi citite de la Yahoo Finance."""


def as_int(value: Any) -> int:
    """Normalizează valorile lipsă/NaN din tabelul pandas Yahoo la zero."""
    try:
        return int(value)
    except (TypeError, ValueError, OverflowError):
        return 0


def get_quote(ticker: str) -> dict[str, Any]:
    """Citește consensul prin yfinance, care gestionează autentificarea Yahoo."""
    stock = yf.Ticker(ticker)
    try:
        targets = stock.get_analyst_price_targets() or {}
        summary = stock.get_recommendations_summary()
    except Exception as exc:  # yfinance ridică tipuri diferite la erori Yahoo
        raise YahooDataError(f"Nu am putut citi datele pentru «{ticker}»: {exc}") from exc

    # Prețul/numele sunt informații auxiliare; lipsa lor nu trebuie să ascundă
    # recomandările care au fost deja preluate cu succes.
    try:
        current_price = stock.fast_info.get("last_price") or targets.get("current")
    except Exception:
        current_price = targets.get("current")
    try:
        name = stock.get_info().get("longName", ticker)
    except Exception:
        name = ticker
    try:
        recent_ratings = stock.get_recommendations()
    except Exception:
        recent_ratings = None

    if not targets and (summary is None or summary.empty):
        raise YahooDataError(f"Nu am găsit date de analiști pentru «{ticker}».")

    counts = {key: 0 for key in ("strongBuy", "buy", "hold", "sell", "strongSell")}
    if summary is not None and not summary.empty:
        latest = summary.iloc[0]
        for key in counts:
            value = latest.get(key, 0)
            counts[key] = as_int(value)

    # Yahoo expune de regulă 0m, -1m, -2m și -3m. Păstrăm aceste rânduri
    # pentru a reproduce tabelul/graficul Analyst Recommendations.
    recommendation_history: list[dict[str, Any]] = []
    if summary is not None and not summary.empty:
        for _, row in summary.head(4).iterrows():
            period = row.get("period", row.name)
            recommendation_history.append({
                "period": str(period),
                **{key: as_int(row.get(key, 0)) for key in counts},
            })

    fallback_ratings: list[dict[str, str]] = []
    if recent_ratings is not None and not recent_ratings.empty:
        # `get_recommendations()` expune datele Yahoo: firm, toGrade,
        # fromGrade, action și data. Coloanele pot varia ușor între versiuni.
        for index, row in recent_ratings.tail(8).iloc[::-1].iterrows():
            date = getattr(index, "strftime", lambda _format: str(index))("%d.%m.%Y")
            fallback_ratings.append({
                "date": date,
                "analyst": str(row.get("Firm", row.get("firm", "—")) or "—"),
                "rating_action": str(row.get("Action", row.get("action", "—")) or "—"),
                "rating": str(row.get("ToGrade", row.get("toGrade", "—")) or "—"),
                "price_action": "—",
                "price_target": "—",
            })

    # Pagina Analysts are date mai bogate decât endpoint-ul standard yfinance:
    # în special Price Action și Price Target. Folosim acea sursă cu prioritate.
    try:
        latest_ratings = find_latest_ratings_in_page(ticker)
    except requests.RequestException:
        latest_ratings = []
    if not latest_ratings:
        latest_ratings = fallback_ratings

    opinions = sum(counts.values())
    weighted = (counts["strongBuy"] + 2 * counts["buy"] + 3 * counts["hold"]
                + 4 * counts["sell"] + 5 * counts["strongSell"])
    average = weighted / opinions if opinions else None
    positive = counts["strongBuy"] + counts["buy"]
    negative = counts["sell"] + counts["strongSell"]
    if opinions == 0:
        recommendation = "n/a"
    elif positive >= negative * 2 and positive > counts["hold"]:
        recommendation = "buy"
    elif negative >= positive * 2 and negative > counts["hold"]:
        recommendation = "sell"
    else:
        recommendation = "hold"

    return {
        "longName": name,
        "regularMarketPrice": current_price,
        "targetMeanPrice": targets.get("mean"),
        "targetLowPrice": targets.get("low"),
        "targetHighPrice": targets.get("high"),
        "numberOfAnalystOpinions": opinions or None,
        "averageAnalystRating": f"{average:.2f} / 5" if average else "n/a",
        "recommendationKey": recommendation,
        "recommendationHistory": recommendation_history,
        "latestRatings": latest_ratings,
    }


def _extract_balanced_json(text: str, start: int) -> str | None:
    """Extrage un obiect JSON echilibrat pornind de la prima acoladă."""
    first = text.find("{", start)
    if first == -1:
        return None
    depth, in_string, escape = 0, False, False
    for index in range(first, len(text)):
        char = text[index]
        if in_string:
            if escape:
                escape = False
            elif char == "\\\\":
                escape = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
        elif char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return text[first : index + 1]
    return None


def find_top_analysts_in_page(ticker: str) -> list[dict[str, Any]]:
    """Încearcă să citească lista Top Analysts din pagina publică Yahoo.

    Yahoo își poate modifica interfața; de aceea codul caută blocuri JSON care
    conțin numele secțiunii și normalizează doar câmpurile utile pentru raport.
    """
    page = get_analysts_page(ticker)
    candidates: list[dict[str, Any]] = []
    for match in re.finditer(r"top analysts|topAnalysts|top_analysts", page, flags=re.IGNORECASE):
        raw = _extract_balanced_json(page, match.start())
        if raw:
            try:
                candidates.append(json.loads(raw))
            except json.JSONDecodeError:
                pass

    def walk(value: Any) -> list[dict[str, Any]]:
        found: list[dict[str, Any]] = []
        if isinstance(value, dict):
            lowered = {str(k).lower(): v for k, v in value.items()}
            name = lowered.get("analystname") or lowered.get("analyst_name") or lowered.get("name")
            firm = lowered.get("firm") or lowered.get("company") or lowered.get("brokerage")
            rating = lowered.get("rating") or lowered.get("recommendation") or lowered.get("action")
            target = lowered.get("price_target") or lowered.get("pricetarget") or lowered.get("targetprice")
            if name and (firm or rating or target):
                found.append({"name": name, "firm": firm, "rating": rating, "target": target})
            for child in value.values():
                found.extend(walk(child))
        elif isinstance(value, list):
            for child in value:
                found.extend(walk(child))
        return found

    unique: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for candidate in candidates:
        for analyst in walk(candidate):
            key = (str(analyst["name"]), str(analyst.get("firm")))
            if key not in seen:
                seen.add(key)
                unique.append(analyst)
    return unique[:10]


def get_analysts_page(ticker: str) -> str:
    """Descarcă o singură dată pagina Analysts, chiar dacă citim două secțiuni."""
    if ticker not in ANALYSTS_PAGE_CACHE:
        response = SESSION.get(ANALYSTS_URL.format(ticker=ticker), headers=HEADERS, timeout=20)
        response.raise_for_status()
        ANALYSTS_PAGE_CACHE[ticker] = response.text
    return ANALYSTS_PAGE_CACHE[ticker]


def _page_json_payloads(page: str) -> list[dict[str, Any]]:
    """Extrage datele serializate din aplicația web Yahoo, fără selectori fragili HTML."""
    payloads: list[dict[str, Any]] = []
    for marker in ("root.App.main", "__NEXT_DATA__", "Latest Ratings", "latestRatings"):
        start = page.find(marker)
        if start == -1:
            continue
        raw = _extract_balanced_json(page, start)
        if raw:
            try:
                data = json.loads(raw)
                if isinstance(data, dict):
                    payloads.append(data)
            except json.JSONDecodeError:
                continue
    return payloads


def _date_for_display(value: Any) -> str:
    if isinstance(value, (int, float)):
        # Unele răspunsuri Yahoo transmit timestamp în secunde sau milisecunde.
        timestamp = value / 1000 if value > 10_000_000_000 else value
        try:
            return datetime.fromtimestamp(timestamp, timezone.utc).strftime("%d.%m.%Y")
        except (OSError, OverflowError, ValueError):
            pass
    return str(value or "—")


def find_latest_ratings_in_page(ticker: str) -> list[dict[str, str]]:
    """Caută dinamic cardurile Yahoo «Latest Ratings» și normalizează câmpurile."""
    page = get_analysts_page(ticker)
    # În pagina actuală Yahoo, cel mai recent rating este redat ca text în
    # cardul „Latest Rating”. Această cale citește chiar structura vizibilă,
    # înaintea blocurilor JSON interne (care se schimbă mai des).
    visible_text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html.unescape(page))).strip()
    card_match = re.search(
        r"Latest Rating\s+Date\s+(?P<date>.+?)\s+Analyst\s+(?P<analyst>.+?)"
        r"\s+Rating Action\s+(?P<rating_action>.+?)\s+Rating\s+(?P<rating>.+?)"
        r"\s+Price Action\s+(?P<price_action>.+?)\s+Price Target\s+(?P<price_target>.+?)"
        r"(?=\s+(?:Statistics|Compare|People Also Watch|Copyright|Terms)\b|$)",
        visible_text,
        flags=re.IGNORECASE,
    )
    if card_match:
        fields = {key: value.strip() for key, value in card_match.groupdict().items()}
        return [{
            "date": fields["date"],
            "analyst": fields["analyst"],
            "rating_action": fields["rating_action"],
            "rating": fields["rating"],
            "price_action": fields["price_action"],
            "price_target": fields["price_target"],
        }]

    records: list[dict[str, str]] = []

    def value_of(data: dict[str, Any], *names: str) -> Any:
        normalized = {str(key).replace("_", "").lower(): value for key, value in data.items()}
        for name in names:
            if name.replace("_", "").lower() in normalized:
                return normalized[name.replace("_", "").lower()]
        return None

    def walk(value: Any) -> None:
        if isinstance(value, dict):
            analyst = value_of(value, "analyst", "analystName", "firm", "company")
            rating = value_of(value, "rating", "toGrade", "grade")
            rating_action = value_of(value, "ratingAction", "action", "rating_action")
            price_action = value_of(value, "priceAction", "price_action")
            price_target = value_of(value, "priceTarget", "targetPrice", "price_target")
            date = value_of(value, "date", "ratingDate", "publishedDate", "epochGradeDate")
            # Evită obiectele generice; un rating Yahoo are un analist/firma și o notă.
            if analyst and rating and (rating_action or price_action or price_target):
                records.append({
                    "date": _date_for_display(date),
                    "analyst": str(analyst),
                    "rating_action": str(rating_action or "—"),
                    "rating": str(rating),
                    "price_action": str(price_action or "—"),
                    "price_target": format_number(price_target) if isinstance(price_target, (int, float)) else str(price_target or "—"),
                })
            for child in value.values():
                walk(child)
        elif isinstance(value, list):
            for child in value:
                walk(child)

    for payload in _page_json_payloads(page):
        walk(payload)
    unique: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()
    for item in records:
        key = (item["date"], item["analyst"], item["rating"], item["price_target"])
        if key not in seen:
            seen.add(key)
            unique.append(item)
    return unique[:8]


def format_number(value: Any) -> str:
    try:
        return f"{float(value):,.2f}"
    except (TypeError, ValueError):
        return "n/a"


def period_label(period: str) -> str:
    """Transformă etichetele Yahoo (0m/-1m/...) în texte ușor de citit."""
    labels = {"0m": "Acum", "-1m": "Acum 1 lună", "-2m": "Acum 2 luni", "-3m": "Acum 3 luni"}
    return labels.get(period, period)


def conclusion(quote: dict[str, Any]) -> str:
    rating = quote.get("averageAnalystRating", "n/a")
    current = quote.get("regularMarketPrice")
    target = quote.get("targetMeanPrice")
    recommendation = str(quote.get("recommendationKey", "n/a")).replace("_", " ").upper()
    if isinstance(current, (int, float)) and isinstance(target, (int, float)) and current > 0:
        upside = (target / current - 1) * 100
        direction = "potențial de creștere" if upside >= 0 else "potențial de scădere"
        return (f"Consensul este {recommendation} (rating mediu: {rating}); ținta medie de "
                f"preț sugerează un {direction} de {abs(upside):.1f}% față de prețul curent.")
    return f"Consensul disponibil este {recommendation} (rating mediu: {rating})."


def upside_percent(quote: dict[str, Any]) -> float | None:
    current, target = quote.get("regularMarketPrice"), quote.get("targetMeanPrice")
    if isinstance(current, (int, float)) and isinstance(target, (int, float)) and current > 0:
        return (target / current - 1) * 100
    return None


def recommendation_style(recommendation: str) -> str:
    return {"buy": "bold green", "hold": "bold yellow", "sell": "bold red"}.get(recommendation, "bold white")


def price_scale(quote: dict[str, Any], width: int = 42) -> Text | None:
    """Construiește o scară Low–Current–Average–High similară Yahoo Finance."""
    low, current = quote.get("targetLowPrice"), quote.get("regularMarketPrice")
    average, high = quote.get("targetMeanPrice"), quote.get("targetHighPrice")
    if not all(isinstance(value, (int, float)) for value in (low, current, average, high)) or high <= low:
        return None
    def position(value: float) -> int:
        return max(0, min(width - 1, round((value - low) / (high - low) * (width - 1))))

    markers = {position(current): ("●", "bold cyan"), position(average): ("◆", "bold green")}
    scale = Text()
    for index in range(width):
        symbol, style = markers.get(index, ("━", "dim blue"))
        scale.append(symbol, style=style)
    return scale


def print_report(ticker: str) -> dict[str, Any]:
    quote = get_quote(ticker)
    quote["ticker"] = ticker
    recommendation = str(quote.get("recommendationKey", "n/a"))
    upside = upside_percent(quote)
    title = Text.assemble((f"{ticker}  ", "bold cyan"), (str(quote.get("longName", ticker)), "white"))
    console.print(Panel(title, subtitle=f"Generat: {datetime.now(timezone.utc).strftime('%d.%m.%Y, %H:%M UTC')}",
                        border_style="bright_blue", expand=False))

    overview = Table(box=box.ROUNDED, show_header=False, expand=False, border_style="blue")
    overview.add_column("Indicator", style="dim", min_width=23)
    overview.add_column("Valoare", min_width=20)
    overview.add_row("Recomandare consens", Text(recommendation.upper(), style=recommendation_style(recommendation)))
    overview.add_row("Rating mediu (1–5)", str(quote.get("averageAnalystRating", "n/a")))
    overview.add_row("Analiști activi", str(quote.get("numberOfAnalystOpinions", "n/a")))
    overview.add_row("Preț curent", format_number(quote.get("regularMarketPrice")))
    overview.add_row("Țintă medie", format_number(quote.get("targetMeanPrice")))
    overview.add_row("Interval ținte", f"{format_number(quote.get('targetLowPrice'))} – {format_number(quote.get('targetHighPrice'))}")
    upside_text = "n/a" if upside is None else f"{upside:+.1f}%"
    overview.add_row("Diferență până la ținta medie", Text(upside_text, style="green" if upside is not None and upside >= 0 else "red"))
    console.print(overview)

    scale = price_scale(quote)
    if scale:
        price_panel = Table.grid(padding=(0, 1))
        price_panel.add_column(justify="left")
        price_panel.add_column(justify="center")
        price_panel.add_column(justify="right")
        price_panel.add_row(f"Low\n{format_number(quote.get('targetLowPrice'))}", scale,
                            f"High\n{format_number(quote.get('targetHighPrice'))}")
        price_panel.add_row("", Text.assemble(("● ", "bold cyan"), (f"Current {format_number(quote.get('regularMarketPrice'))}   ", "cyan"),
                                                 ("◆ ", "bold green"), (f"Average {format_number(quote.get('targetMeanPrice'))}", "green")), "")
        console.print(Panel(price_panel, title="Price Targets — Yahoo Finance", border_style="blue", expand=False))

    history = quote.get("recommendationHistory", [])
    if history:
        trend = Table(title="Analyst Recommendations — ultimele 4 luni", box=box.ROUNDED, border_style="magenta")
        trend.add_column("Perioadă", style="bold")
        for column in ("Strong Buy", "Buy", "Hold", "Underperform", "Sell", "Total"):
            trend.add_column(column, justify="right")
        for row in history:
            total = row["strongBuy"] + row["buy"] + row["hold"] + row["sell"] + row["strongSell"]
            trend.add_row(period_label(row["period"]), str(row["strongBuy"]), str(row["buy"]), str(row["hold"]),
                          str(row["sell"]), str(row["strongSell"]), f"[bold]{total}[/bold]")
        console.print(trend)
    else:
        console.print("[yellow]Yahoo nu a returnat istoricul recomandărilor pentru acest simbol.[/yellow]")

    latest_ratings = quote.get("latestRatings", [])
    if latest_ratings:
        ratings_table = Table(title="Latest Ratings", box=box.ROUNDED, border_style="yellow")
        ratings_table.add_column("Dată", style="dim")
        ratings_table.add_column("Analyst", style="bold")
        ratings_table.add_column("Rating Action")
        ratings_table.add_column("Rating")
        ratings_table.add_column("Price Action")
        ratings_table.add_column("Price Target", justify="right")
        for item in latest_ratings:
            ratings_table.add_row(item["date"], item["analyst"], item["rating_action"], item["rating"],
                                  item["price_action"], item["price_target"])
        console.print(ratings_table)
    else:
        console.print("[yellow]Yahoo nu a returnat ultimele ratinguri pentru acest simbol.[/yellow]")

    try:
        analysts = find_top_analysts_in_page(ticker)
    except requests.RequestException as exc:
        analysts = []

    if analysts:
        top_table = Table(title="Top Analysts — pagina Yahoo Finance", box=box.ROUNDED, border_style="cyan")
        top_table.add_column("Analist", style="bold")
        top_table.add_column("Firmă")
        top_table.add_column("Recomandare")
        top_table.add_column("Țintă", justify="right")
        for item in analysts:
            top_table.add_row(str(item["name"]), str(item.get("firm") or "—"), str(item.get("rating") or "—"),
                              str(item.get("target") or "—"))
        console.print(top_table)
    else:
        console.print("[yellow]Top Analysts nu a fost furnizat în format structurat de Yahoo în această sesiune.[/yellow]")

    console.print(Panel(conclusion(quote), title="Concluzie", border_style=recommendation_style(recommendation)))
    return quote


def read_tickers(values: list[str], file_path: str | None) -> list[str]:
    """Acceptă AAPL MSFT, AAPL,MSFT sau un fișier cu câte un simbol pe rând."""
    raw = list(values)
    if file_path:
        try:
            with open(file_path, encoding="utf-8") as handle:
                raw.extend(line.split("#", 1)[0].strip() for line in handle)
        except OSError as exc:
            raise YahooDataError(f"Nu pot citi fișierul cu simboluri: {exc}") from exc
    tickers: list[str] = []
    for value in raw:
        for ticker in value.split(","):
            ticker = ticker.strip().upper()
            if ticker and ticker not in tickers:
                tickers.append(ticker)
    return tickers


def print_comparison(reports: list[dict[str, Any]]) -> None:
    if len(reports) < 2:
        return
    table = Table(title="Comparație rapidă", box=box.DOUBLE_EDGE, border_style="green")
    table.add_column("Simbol", style="bold cyan")
    table.add_column("Consens")
    table.add_column("Preț", justify="right")
    table.add_column("Țintă medie", justify="right")
    table.add_column("Potențial", justify="right")
    table.add_column("Analiști", justify="right")
    for report in reports:
        upside = upside_percent(report)
        potential = "n/a" if upside is None else f"{upside:+.1f}%"
        recommendation = str(report.get("recommendationKey", "n/a"))
        table.add_row(report["ticker"], Text(recommendation.upper(), style=recommendation_style(recommendation)),
                      format_number(report.get("regularMarketPrice")), format_number(report.get("targetMeanPrice")),
                      Text(potential, style="green" if upside is not None and upside >= 0 else "red"),
                      str(report.get("numberOfAnalystOpinions") or "n/a"))
    console.print(table)


def safe_sheet_name(ticker: str, used: set[str]) -> str:
    """Creează un nume de foaie Excel valid și unic."""
    base = re.sub(r"[\\/*?:\[\]]", "_", ticker)[:31] or "Companie"
    name, number = base, 2
    while name in used:
        suffix = f"_{number}"
        name = f"{base[:31 - len(suffix)]}{suffix}"
        number += 1
    used.add(name)
    return name


def style_header(cells: list[Any], fill: str = "1F4E78") -> None:
    for cell in cells:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center")


def write_excel_report(reports: list[dict[str, Any]], tickers: list[str]) -> Path | None:
    """Salvează raportul Excel în ./Output.

    Foaia Sumar conține câte un rând pentru fiecare ticker și toate valorile
    curente (0m): Strong Buy, Buy, Hold, Underperform și Sell. Workbook-ul
    păstrează în plus câte o foaie detaliată pentru fiecare companie.
    """
    if not reports:
        return None

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Sumar"

    summary_headers = [
        "Ticker",
        "Companie",
        "Consens",
        "Strong Buy",
        "Buy",
        "Hold",
        "Underperform",
        "Sell",
        "Strong Buy %",
        "Buy %",
        "Hold %",
        "Sell %",
        "Total Analysts",
        "Average Analyst Rating",
        "Preț curent",
        "Țintă medie",
        "Țintă minimă",
        "Țintă maximă",
        "Potențial",
        "Ranking Score",
    ]
    last_summary_column = get_column_letter(len(summary_headers))

    summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(summary_headers))
    summary["A1"] = "Raport analiști Yahoo Finance"
    summary["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    summary["A1"].alignment = Alignment(horizontal="center")
    summary["A2"] = "Data interogării"
    summary["B2"] = datetime.now().astimezone().strftime("%d.%m.%Y %H:%M")
    summary["A3"] = "Sursă"
    summary["B3"] = "Yahoo Finance"
    summary.append([])
    summary.append(summary_headers)
    style_header(list(summary[5]))

    for report in reports:
        upside = upside_percent(report)
        history = report.get("recommendationHistory", [])
        current = history[0] if history else {}

        strong_buy = as_int(current.get("strongBuy", 0))
        buy = as_int(current.get("buy", 0))
        hold = as_int(current.get("hold", 0))
        underperform = as_int(current.get("sell", 0))
        sell = as_int(current.get("strongSell", 0))
        total = strong_buy + buy + hold + underperform + sell

        strong_buy_pct = strong_buy / total if total else None
        buy_pct = buy / total if total else None
        hold_pct = hold / total if total else None
        sell_pct = (underperform + sell) / total if total else None

        # Scor orientativ: pune accent pe recomandările pozitive și penalizează
        # Underperform/Sell. Nu reprezintă o recomandare de investiție.
        ranking_score = (
            strong_buy * 5
            + buy * 3
            + hold
            - underperform * 2
            - sell * 4
        )

        summary.append([
            report["ticker"],
            report.get("longName", "—"),
            str(report.get("recommendationKey", "n/a")).upper(),
            strong_buy,
            buy,
            hold,
            underperform,
            sell,
            strong_buy_pct,
            buy_pct,
            hold_pct,
            sell_pct,
            total,
            report.get("averageAnalystRating", "n/a"),
            report.get("regularMarketPrice"),
            report.get("targetMeanPrice"),
            report.get("targetLowPrice"),
            report.get("targetHighPrice"),
            upside / 100 if upside is not None else None,
            ranking_score,
        ])

    for row in summary.iter_rows(min_row=6, max_row=summary.max_row):
        for index in (8, 9, 10, 11):
            row[index].number_format = "0.0%"
        for index in (14, 15, 16, 17):
            row[index].number_format = "#,##0.00"
        row[18].number_format = "0.0%"

    widths = (
        12, 35, 14, 12, 10, 10, 15, 10, 14, 12,
        12, 12, 15, 22, 15, 15, 15, 15, 13, 15,
    )
    for column, width in enumerate(widths, start=1):
        summary.column_dimensions[get_column_letter(column)].width = width

    summary.auto_filter.ref = f"A5:{last_summary_column}{summary.max_row}"
    summary.freeze_panes = "A6"

    used_names = {"Sumar"}
    for report in reports:
        sheet = workbook.create_sheet(safe_sheet_name(report["ticker"], used_names))
        sheet.merge_cells("A1:G1")
        sheet["A1"] = f"{report.get('longName', report['ticker'])} ({report['ticker']})"
        sheet["A1"].font = Font(size=15, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        sheet["A1"].alignment = Alignment(horizontal="center")

        sheet["A3"] = "Analyst Recommendation"
        sheet["A3"].font = Font(size=12, bold=True, color="1F1F1F")
        overview = [
            ("Recomandare consens", str(report.get("recommendationKey", "n/a")).upper()),
            ("Rating mediu", report.get("averageAnalystRating", "n/a")),
            ("Analiști activi", report.get("numberOfAnalystOpinions", "n/a")),
            ("Preț curent", report.get("regularMarketPrice")),
            ("Țintă medie", report.get("targetMeanPrice")),
            ("Țintă minimă", report.get("targetLowPrice")),
            ("Țintă maximă", report.get("targetHighPrice")),
            (
                "Potențial până la ținta medie",
                upside_percent(report) / 100 if upside_percent(report) is not None else None,
            ),
        ]
        for row_number, (label, value) in enumerate(overview, start=4):
            sheet.cell(row_number, 1, label).font = Font(bold=True)
            sheet.cell(row_number, 2, value)
        for row_number in range(7, 11):
            sheet.cell(row_number, 2).number_format = "#,##0.00"
        sheet["B11"].number_format = "0.0%"

        sheet["A13"] = "Analyst Recommendations — ultimele 4 luni"
        sheet["A13"].font = Font(size=12, bold=True)
        history_headers = [
            "Perioadă", "Strong Buy", "Buy", "Hold",
            "Underperform", "Sell", "Total",
        ]
        for column, value in enumerate(history_headers, start=1):
            sheet.cell(14, column, value)
        style_header(list(sheet[14]))

        for row in report.get("recommendationHistory", []):
            total = sum(row[key] for key in ("strongBuy", "buy", "hold", "sell", "strongSell"))
            sheet.append([
                period_label(row["period"]),
                row["strongBuy"],
                row["buy"],
                row["hold"],
                row["sell"],
                row["strongSell"],
                total,
            ])

        ratings_start = max(20, sheet.max_row + 2)
        sheet.cell(ratings_start, 1, "Latest Rating")
        sheet.cell(ratings_start, 1).font = Font(size=12, bold=True)
        rating_headers = [
            "Date", "Analyst", "Rating Action", "Rating",
            "Price Action", "Price Target",
        ]
        for column, value in enumerate(rating_headers, start=1):
            sheet.cell(ratings_start + 1, column, value)
        style_header(list(sheet[ratings_start + 1]))

        for item in report.get("latestRatings", []):
            sheet.append([
                item["date"],
                item["analyst"],
                item["rating_action"],
                item["rating"],
                item["price_action"],
                item["price_target"],
            ])

        conclusion_row = max(sheet.max_row + 2, ratings_start + 3)
        sheet.cell(conclusion_row, 1, "Concluzie").font = Font(size=12, bold=True)
        sheet.merge_cells(
            start_row=conclusion_row + 1,
            start_column=1,
            end_row=conclusion_row + 2,
            end_column=6,
        )
        conclusion_cell = sheet.cell(conclusion_row + 1, 1, conclusion(report))
        conclusion_cell.alignment = Alignment(wrap_text=True, vertical="top")

        for column, width in enumerate((19, 24, 20, 20, 20, 18, 12), start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = "A4"

    output_dir = Path.cwd() / "Output"
    output_dir.mkdir(parents=True, exist_ok=True)
    date_label = datetime.now().astimezone().strftime("%d.%m.%y")
    if len(tickers) > 3:
        label = f"MULTI_{len(tickers)}_TICKERS"
    else:
        label = " ".join(tickers)[:120]
    target = output_dir / f"{date_label} {label}.xlsx"
    number = 2
    while target.exists():
        target = output_dir / f"{date_label} {label} ({number}).xlsx"
        number += 1
    workbook.save(target)
    return target

def main() -> None:
    parser = argparse.ArgumentParser(description="Raport de analiști din Yahoo Finance")
    parser.add_argument("tickers", nargs="*", help="simboluri: AAPL MSFT NVDA sau AAPL,MSFT,NVDA")
    parser.add_argument("-f", "--file", help="fișier text cu simboluri; un simbol pe rând")
    args = parser.parse_args()
    tickers = read_tickers(args.tickers, args.file)
    if not tickers:
        parser.error("introdu cel puțin un simbol sau folosește --file lista.txt")
    reports: list[dict[str, Any]] = []
    console.print(Panel.fit("[bold]Yahoo Finance · Dashboard de analiză a analiștilor[/bold]", border_style="blue"))
    for ticker in tickers:
        try:
            reports.append(print_report(ticker))
        except (requests.RequestException, ValueError, YahooDataError) as exc:
            console.print(f"[bold red]Eroare pentru {ticker}:[/bold red] {exc}")
    print_comparison(reports)
    try:
        workbook_path = write_excel_report(reports, tickers)
        if workbook_path:
            console.print(f"[bold green]Fișier Excel creat:[/bold green] {workbook_path}")
    except Exception as exc:
        console.print(f"[yellow]Nu am putut crea fișierul Excel: {exc}[/yellow]")
    console.print("[dim]Date informative din Yahoo Finance; nu reprezintă consultanță financiară.[/dim]")


if __name__ == "__main__":
    main()
